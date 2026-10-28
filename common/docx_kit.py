"""Построитель отчётов Word (python-docx) с тем же API, что и PDF-отчёт."""

from __future__ import annotations

import datetime as _dt
import os
from typing import Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

COLOR_ACCENT = RGBColor(0x1F, 0x5F, 0xA9)
COLOR_INK = RGBColor(0x1A, 0x20, 0x2C)
COLOR_MUTED = RGBColor(0x4A, 0x55, 0x68)
CELL_HEAD = "E8EEF6"
CELL_TOTAL = "F2F5F9"


def _shade(cell, hex_color: str) -> None:
    """Заливка ячейки таблицы (python-docx не имеет прямого API)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _add_page_number(paragraph) -> None:
    """Вставляет поле PAGE в параграф (нумерация страниц в колонтитуле)."""
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


class ReportDocx:
    """Документ Word: разделы, таблицы, рисунки, примечания."""

    def __init__(self, title: str, program: str, subtitle: str = "", doc_no: str = "") -> None:
        self.title = title
        self.program = program
        self.subtitle = subtitle
        self.doc_no = doc_no or f"РР-{_dt.date.today():%Y%m%d}"
        self.doc = Document()
        self._section = 0
        self._figure_no = 0
        self._table_no = 0
        self._setup()

    # ------------------------------------------------------------------ настройка
    def _setup(self) -> None:
        doc = self.doc
        section = doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(9.5)
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = 1.08

        for name, size, color in (("Heading 1", 13, COLOR_ACCENT), ("Heading 2", 10.5, COLOR_INK)):
            style = doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(9)
            style.paragraph_format.space_after = Pt(3)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"{self.program} · {self.doc_no} · стр. ")
        run.font.size = Pt(7.5)
        run.font.color.rgb = COLOR_MUTED
        _add_page_number(footer)

        head = doc.add_paragraph()
        r = head.add_run(self.title)
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_INK
        head.paragraph_format.space_after = Pt(1)
        if self.subtitle:
            sub = doc.add_paragraph()
            rs = sub.add_run(self.subtitle)
            rs.font.size = Pt(9.5)
            rs.font.color.rgb = COLOR_MUTED
        meta = doc.add_paragraph()
        rm = meta.add_run(f"Документ {self.doc_no} · {_dt.date.today():%d.%m.%Y}")
        rm.font.size = Pt(8)
        rm.font.color.rgb = COLOR_MUTED
        meta.paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------------- разметка
    def h1(self, text: str) -> None:
        self._section += 1
        self.doc.add_heading(f"{self._section}. {text}", level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def p(self, text: str, small: bool = False) -> None:
        par = self.doc.add_paragraph(text)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if small or not text:
            for run in par.runs:
                run.font.size = Pt(8)

    def bullet(self, items: Sequence[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def note(self, text: str) -> None:
        par = self.doc.add_paragraph(text)
        for run in par.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = COLOR_MUTED
        par.paragraph_format.left_indent = Cm(0.3)


    # -------------------------------------------------------------------- таблицы
    def table(
        self,
        header: Sequence[str],
        rows: Sequence[Sequence[str]],
        widths_mm: Sequence[float] | None = None,
        aligns: Sequence[str] | None = None,
        caption: str = "",
        total_last: bool = False,
    ) -> None:
        if caption:
            self._table_no += 1
            cap = self.doc.add_paragraph()
            run = cap.add_run(f"Таблица {self._table_no}. {caption}")
            run.font.size = Pt(8)
            run.font.color.rgb = COLOR_MUTED
            cap.paragraph_format.space_after = Pt(2)

        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        aligns = list(aligns or ["l"] * len(header))
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                "r": WD_ALIGN_PARAGRAPH.RIGHT}

        head_cells = table.rows[0].cells
        for i, text in enumerate(header):
            head_cells[i].text = ""
            par = head_cells[i].paragraphs[0]
            par.alignment = amap.get(aligns[i], WD_ALIGN_PARAGRAPH.LEFT)
            run = par.add_run(str(text))
            run.bold = True
            run.font.size = Pt(8.5)
            _shade(head_cells[i], CELL_HEAD)

        for r_index, row in enumerate(rows):
            cells = table.add_row().cells
            is_total = total_last and r_index == len(rows) - 1
            for i, value in enumerate(row):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                par.alignment = amap.get(aligns[i], WD_ALIGN_PARAGRAPH.LEFT)
                par.paragraph_format.space_after = Pt(0)
                run = par.add_run(str(value))
                run.font.size = Pt(8.5)
                run.bold = is_total
                if is_total:
                    _shade(cells[i], CELL_TOTAL)

        if widths_mm:
            for row in table.rows:
                for i, width in enumerate(widths_mm):
                    row.cells[i].width = Mm(width)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def params(self, rows: Sequence[tuple[str, str, str]], caption: str = "") -> None:
        self.table(["Параметр", "Обозн.", "Значение"], list(rows),
                   widths_mm=[88, 26, 61], aligns=["l", "c", "r"], caption=caption)

    def figure(self, png_path: str, caption: str = "", width_mm: float = 172) -> None:
        self.doc.add_picture(png_path, width=Mm(width_mm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            self._figure_no += 1
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(f"Рисунок {self._figure_no}. {caption}")
            run.font.size = Pt(8)
            run.font.color.rgb = COLOR_MUTED

    # --------------------------------------------------------------------- вывод
    def save(self, path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.doc.save(path)
        return path
